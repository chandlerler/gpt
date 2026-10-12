"""Browser integration with explicit Chrome API/platform fixtures, not a live-site test."""
import asyncio
import base64
import json
import mimetypes
import shutil
import struct
import zipfile
import zlib
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from pathlib import Path
from datetime import datetime, timezone
from playwright.async_api import async_playwright

ROOT = Path(__file__).resolve().parents[1]
EXT = ROOT / 'extension'
RESULTS = ROOT / 'test-results'

def png():
    def chunk(t, d):
        return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t+d))
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 100, 100, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress((b'\0'+b'\x22\x55\x88'*100)*100)) + chunk(b'IEND',b'')

FIXTURE = r'''(() => {
  const event=()=>{const listeners=[];return {addListener:f=>listeners.push(f),removeListener:f=>{const i=listeners.indexOf(f);if(i>=0)listeners.splice(i,1);},emit:(...args)=>listeners.forEach(f=>f(...args))};};
  const data=Object.assign({aisFirstUseAt:1,ais_license_cache:{expiresAt:1,features:[]},ais_license_code:'expired-test-code'},JSON.parse(localStorage.getItem('fixture-store')||'{}'));
  const changed=event();
  const answer=(v,cb)=>{if(cb)queueMicrotask(()=>cb(v));return Promise.resolve(v);};
  window.fixtureMessages=[];
  const storage={get(keys,cb){let r;if(keys==null)r={...data};else if(typeof keys==='object'&&!Array.isArray(keys))r=Object.fromEntries(Object.entries(keys).map(([k,v])=>[k,data[k]??v]));else r=Object.fromEntries((Array.isArray(keys)?keys:[keys]).filter(k=>k in data).map(k=>[k,data[k]]));return answer(r,cb);},set(obj,cb){Object.assign(data,obj);localStorage.setItem('fixture-store',JSON.stringify(data));changed.emit(Object.fromEntries(Object.entries(obj).map(([k,v])=>[k,{newValue:v}])),'local');return answer(undefined,cb);},remove(keys,cb){for(const k of Array.isArray(keys)?keys:[keys])delete data[k];return answer(undefined,cb);}};
  window.chrome={runtime:{id:'local-test',getManifest:()=>window.fixtureManifest,getURL:f=>'https://fixture.test/extension/'+f,onMessage:event(),sendMessage(m,cb){fixtureMessages.push(m);if(m.action==='licenseFetch')throw Error('Unexpected license request');let r={ok:true,success:true};if(m.action==='aisAnchorRules')r={ok:false};if(m.action==='getHistory')r={history:[]};if(m.action==='mainWorldPing')r={ok:true};if(m.action==='mainWorldGet1688RawData'||m.action==='mainWorldGetPddRawData')r={success:false};if(m.action==='fetchDescImages')r={ok:true,images:[]};if(m.action==='downloadImages')r={ok:true,count:m.images.length,down:m.images.length,failed:0};return answer(r,cb);}},storage:{local:storage,sync:storage,onChanged:changed},i18n:{getUILanguage:()=> 'zh-CN'},downloads:{showDefaultFolder(){}}};
})();'''

IMAGE_TRANSPORT = r'''(() => { const send=chrome.runtime.sendMessage;
chrome.runtime.sendMessage=function(m,cb){
  if(m.action!=='fetchReviewImage')return send(m,cb);
  fixtureMessages.push(m);
  const task=fetch(m.url,{credentials:'include'}).then(async r=>{
    if(!r.ok)throw Error('图片 HTTP '+r.status);
    const bytes=new Uint8Array(await r.arrayBuffer());
    return {ok:true,mime:r.headers.get('content-type'),base64:btoa(Array.from(bytes,b=>String.fromCharCode(b)).join(''))};
  }).catch(e=>({ok:false,error:e.message}));
  if(cb)task.then(cb);return task;
}; })();'''

async def main():
    RESULTS.mkdir(exist_ok=True)
    manifest=json.loads((EXT/'manifest.json').read_text())
    checks=[]
    errors=[]
    network=[]
    async with async_playwright() as p:
        browser=await p.chromium.launch(executable_path=shutil.which('chromium'),headless=True)
        context=await browser.new_context(accept_downloads=True)
        context.set_default_timeout(15000)
        async def route(r):
            u=r.request.url
            if u.startswith('https://fixture.test/extension/'):
                rel=u.split('/extension/',1)[1].split('?',1)[0]
                file=(EXT/rel).resolve()
                assert file.is_relative_to(EXT) and file.is_file(),u
                await r.fulfill(body=file.read_bytes(),content_type=mimetypes.guess_type(str(file))[0] or 'application/octet-stream')
            elif u.startswith('https://img.alicdn.com/'):
                await r.fulfill(status=404 if 'missing' in u else 200,body=png(),content_type='image/png',headers={'Access-Control-Allow-Origin':r.request.headers.get('origin','https://fixture.test'),'Access-Control-Allow-Credentials':'true'})
            elif u.startswith('https://item.taobao.com/'):
                await r.fulfill(body='''<!doctype html><meta charset="utf-8"><title>测试商品</title><body><h1>测试商品</h1><div id="picGalleryEle"><img width="100" height="100" src="https://img.alicdn.com/main.png"></div></body>''',content_type='text/html')
            else:
                network.append(u)
                await r.abort()
        await context.route('**/*',route)
        page=await context.new_page()
        page.on('pageerror',lambda e:errors.append(str(e)))
        await page.add_init_script('window.fixtureManifest='+json.dumps(manifest)+';')
        await page.add_init_script(FIXTURE)
        await page.goto('https://item.taobao.com/item.htm?id=123456789')
        await page.evaluate(IMAGE_TRANSPORT)
        await page.clock.set_system_time(datetime(2040,1,1,tzinfo=timezone.utc))
        for entry in manifest['content_scripts']:
            for file in entry['js']:
                await page.add_script_tag(path=EXT/file)
        await page.wait_for_function('window.AISIMG && AISIMG.bound && window.__AIS_SHADOW__ && __AIS_DOC().getElementById("ais-license-tag")')
        await page.evaluate('updateLicenseTag();refreshMemberCache();')
        status=await page.evaluate('''async () => ({text:__AIS_DOC().getElementById('ais-license-tag').textContent,member:isMemberNow(),batch:await requirePaidFeature('batch_detail'),zip:requireMember('ZIP'),keys:await chrome.storage.local.get(['aisFirstUseAt'])})''')
        assert status['text']=='本地版' and status['member'] and status['batch'] and status['zip'],status
        assert status['keys']['aisFirstUseAt']==1
        checks.append('Full manifest content-script initialization with expired storage and a 2040 clock; local status and action gates succeed.')
        html=await page.evaluate('__AIS_DOC().innerHTML')
        assert 'manfaka.com' not in html and '>会员</span>' not in html
        checks.append('Injected Shadow DOM has no purchase links or member badges.')
        floating=await page.evaluate('''() => {
          const root=__AIS_DOC();root.getElementById('ais-panel').classList.add('ais-minimized');
          return {animation:getComputedStyle(root.querySelector('.ais-float-pulse')).animationName,motion:root.getElementById('ais-float-motion').checked,visible:root.getElementById('ais-float-visible').checked};
        }''')
        assert floating=={'animation':'none','motion':False,'visible':True},floating
        await page.evaluate("__AIS_DOC().getElementById('ais-float-motion').click()")
        assert await page.evaluate("getComputedStyle(__AIS_DOC().querySelector('.ais-float-pulse')).animationName")=='aisFloatPulse'
        await page.evaluate("__AIS_DOC().getElementById('ais-float-visible').click();__AIS_DOC().getElementById('ais-float-btn').style.display='flex'")
        assert await page.evaluate("getComputedStyle(__AIS_DOC().getElementById('ais-float-btn')).display")=='none'
        checks.append('Floating ball animation defaults off; independent switches persist and hiding overrides old inline display writes.')

        # Feed representative collected data into the real panel state, then invoke real button handlers.
        await page.evaluate('''() => {
          AISIMG.itemTitle='测试商品'; AISIMG.itemId='123456789';
          AISIMG.reviews=[{reviewId:'1',content:'质量很差，已经破损',rating:1,time:'2039-12-31',images:['https://img.alicdn.com/review.png'],skuInfo:'颜色:红色'}];
          AISIMG.skuRows=[{skuId:'101',price:'19.9',stock:10,properties:'红色',image:'https://img.alicdn.com/sku.png'}];
          AISIMG.analysis={total:1,negativeCount:1,negativeReviews:[{reviewId:'1',content:'质量很差，已经破损',reasons:['质量'],topics:['质量']}],ratingDist:{1:1},bySku:[],byTime:[],negativeTopics:[]};
          AISIMG.main=['https://img.alicdn.com/main.png']; AISIMG.desc=['https://img.alicdn.com/detail.png']; AISIMG.readFinished=true;
        }''')
        ids=await page.evaluate("Array.from(__AIS_DOC().querySelectorAll('[id]')).map(x=>x.id).filter(x=>/rev.*(csv|json|xlsx)|sku.*(csv|xlsx)/.test(x))")
        assert 'aisimg-revexp-json' in ids,ids
        await page.evaluate("__AIS_DOC().getElementById('aisimg-revexp-json').click()")
        await page.wait_for_function("fixtureMessages.some(m=>m.action==='downloadData'&&m.filename.endsWith('_评论.json'))")
        msg=await page.evaluate("fixtureMessages.find(m=>m.action==='downloadData'&&m.filename.endsWith('_评论.json'))")
        rows=json.loads(base64.b64decode(msg['base64']))
        assert rows[0]['图片地址']=='https://img.alicdn.com/review.png',rows
        checks.append('Actual comment JSON export includes review image URLs without activation.')
        await page.evaluate('''() => {
          AISIMG.reviews[0].images.push('https://img.alicdn.com/second.png');
          AISIMG.reviews[0].appendImages=['https://img.alicdn.com/append.png'];
          AISIMG.reviews.push({reviewId:'2',content:'无图评价',time:'2039-12-30',images:[]});
          AISIMG.reviews.push({reviewId:'3',content:'图片失败样本',time:'2039-12-29',images:['https://img.alicdn.com/review.png','https://img.alicdn.com/missing.png']});
          __AIS_DOC().getElementById('aisimg-revexp-xlsx').click();
        }''')
        await page.wait_for_function("fixtureMessages.some(m=>m.action==='downloadData'&&m.filename.endsWith('_评论.xlsx'))")
        msg=await page.evaluate("fixtureMessages.find(m=>m.action==='downloadData'&&m.filename.endsWith('_评论.xlsx'))")
        file=RESULTS/'comments-with-images.xlsx';file.write_bytes(base64.b64decode(msg['base64']))
        wb=load_workbook(file);ws=wb['评论']
        assert len(ws._images)==4,len(ws._images)
        headers={c.value:c.column for c in ws[3]}
        content_rows={ws.cell(row,headers['内容']).value:row for row in range(4,ws.max_row+1)}
        expected_rows=sorted([content_rows['质量很差，已经破损']-1]*3+[content_rows['图片失败样本']-1])
        assert sorted(i.anchor._from.row for i in ws._images)==expected_rows
        assert ws.row_dimensions[content_rows['质量很差，已经破损']].height==96
        assert 'HTTP 404' in ws.cell(content_rows['图片失败样本'],headers['评论图片2']).value
        assert 'missing.png' in ws.cell(content_rows['图片失败样本'],headers['图片地址']).value
        with zipfile.ZipFile(file) as z:
            assert len([n for n in z.namelist() if n.startswith('xl/media/') and not n.endswith('/')])==3
            for name in z.namelist():
                if name.endswith('.xml') or name.endswith('.rels'):ET.fromstring(z.read(name))
        checks.append('Actual comment Excel embeds multiple and appended images in the correct rows; offline parser reads drawings, repeated images share media and failed images retain links plus error text.')
        await page.evaluate("__AIS_DOC().getElementById('aisimg-revnegexp').click()")
        await page.wait_for_function("fixtureMessages.some(m=>m.action==='downloadData'&&m.filename.endsWith('_差评分析.xlsx'))")
        checks.append('Risk analysis export button reaches the download path without a purchase dialog.')
        await page.evaluate("__AIS_DOC().getElementById('aisimg-revimg').click()")
        await page.wait_for_function("fixtureMessages.some(m=>m.action==='downloadImages'&&m.asZip&&m.images.some(i=>i.url==='https://img.alicdn.com/review.png'))")
        checks.append('Comment image button dispatches real background download request.')
        await page.evaluate("__AIS_DOC().getElementById('aisimg-skuexp-xlsx').click()")
        await page.wait_for_function("fixtureMessages.some(m=>m.action==='downloadData'&&m.filename.endsWith('_SKU.xlsx'))")
        msg=await page.evaluate("fixtureMessages.find(m=>m.action==='downloadData'&&m.filename.endsWith('_SKU.xlsx'))")
        file=RESULTS/'sku.xlsx'
        file.write_bytes(base64.b64decode(msg['base64']))
        with zipfile.ZipFile(file) as z:
            assert 'xl/workbook.xml' in z.namelist()
        checks.append('SKU Excel button produces a structurally valid XLSX workbook.')
        assert not errors,errors

        # Exercise the real offscreen ZIP and image conversion, with HTTP image fixtures.
        off=await context.new_page()
        await off.add_init_script('window.fixtureManifest='+json.dumps(manifest)+';')
        await off.add_init_script(FIXTURE)
        await off.goto('https://fixture.test/extension/background/offscreen.html')
        result=await off.evaluate('''async () => {
          const out=await doRun({asZip:true,itemTitle:'测试商品',images:[{url:'https://img.alicdn.com/review.png',type:'review',folder:'评论图片',format:'image/jpeg',quality:90,scale:0.5}]});
          const blob=await (await fetch(out.url)).blob();
          const bytes=new Uint8Array(await blob.arrayBuffer());
          return {out,base64:btoa(Array.from(bytes,x=>String.fromCharCode(x)).join(''))};
        }''')
        file=RESULTS/'review-images.zip'
        file.write_bytes(base64.b64decode(result['base64']))
        with zipfile.ZipFile(file) as z:
            images=[n for n in z.namelist() if n.endswith('.jpg')]
            assert len(images)==1,z.namelist()
            assert z.read(images[0]).startswith(b'\xff\xd8\xff')
        checks.append('Real offscreen Canvas converts a review image to JPEG and produces a readable ZIP.')
        await page.reload()
        await page.evaluate(IMAGE_TRANSPORT)
        for entry in manifest['content_scripts']:
            for file in entry['js']:
                await page.add_script_tag(path=EXT/file)
        await page.wait_for_function('window.AISIMG && AISIMG.bound')
        prefs=await page.evaluate("({visible:__AIS_DOC().getElementById('ais-float-visible').checked,motion:__AIS_DOC().getElementById('ais-float-motion').checked})")
        assert prefs=={'visible':False,'motion':True},prefs
        await page.evaluate("chrome.runtime.onMessage.emit({action:'togglePanel'}, {}, ()=>{})")
        assert await page.evaluate("!__AIS_DOC().getElementById('ais-panel').classList.contains('ais-minimized')")
        assert await page.evaluate("getComputedStyle(__AIS_DOC().getElementById('ais-float-btn')).display")=='none'
        checks.append('Both floating preferences survive a page reload; the toolbar toggle message still opens the panel with the ball hidden.')
        assert not any('/api/license' in u for u in network),network
        (RESULTS/'browser-report.json').write_text(json.dumps({'checks':checks,'page_errors':errors,'blocked_external_requests':network,'scope':'Chromium with mocked Chrome APIs and fixture platform/media responses; not real extension installation'},ensure_ascii=False,indent=2))
        await browser.close()
    print(json.dumps({'passed':len(checks),'checks':checks},ensure_ascii=False,indent=2))

asyncio.run(main())
